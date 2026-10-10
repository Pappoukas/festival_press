"""
clippings.py — Ανάγνωση αποδελτίωσης (.docx) και βοηθητικές συναρτήσεις κειμένου.

Αναμενόμενη δομή του .docx (όπως στο Publicity5thCBF.docx):
  • Title                      → τίτλος αποδελτίωσης
  • Έντονη παράγραφος «Δ. Τύπου» → κύμα δελτίου τύπου (π.χ. «4ο Δ. Τύπου 8/5 (Παμούκ)»)
  • Heading 1 «ΗΗ.ΜΜ[.ΕΕΕΕ]»   → ημερομηνία για τους συνδέσμους που ακολουθούν
  • Heading 1 «Συνεντεύξεις»   → οι επόμενοι σύνδεσμοι σημειώνονται ως συνεντεύξεις
  • Heading 1 «…domain…»       → σύνοψη, αγνοείται
  • Παράγραφοι με URL          → δημοσιεύματα (το κείμενο μετά το URL κρατιέται ως σημείωση)
"""
from __future__ import annotations

import re
import unicodedata
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlparse, urlunparse

import pandas as pd
from docx import Document

URL_RE = re.compile(r"https?://[^\s<>\"\)\]]+")
DATE_HEAD_RE = re.compile(r"^\s*(\d{1,2})\.(\d{1,2})(?:\.(\d{4}))?\s*$")
DM_RE = re.compile(r"(\d{1,2})\s*/\s*(\d{1,2})")
URL_DATE_RE = re.compile(r"/(20\d{2})/(\d{1,2})/(\d{1,2})/")
TRACKING = ("fbclid", "gclid", "si", "igshid", "mc_cid", "mc_eid")

# Ένα κύμα Δ.Τ. «καλύπτει» δημοσιεύματα έως τόσες ημέρες μετά την ημερομηνία του
WAVE_WINDOW_DAYS = 20
NO_WAVE = "Χωρίς συγκεκριμένο Δ.Τ."


# ─────────────────────────────────────────────
# Κείμενο
# ─────────────────────────────────────────────
def normalize(text: str) -> str:
    """Πεζά, χωρίς τόνους, τελικό σίγμα → σ. Για αντιστοίχιση λέξεων-κλειδιών."""
    if not isinstance(text, str):
        return ""
    t = unicodedata.normalize("NFD", text.lower())
    t = "".join(ch for ch in t if unicodedata.category(ch) != "Mn")
    return t.replace("ς", "σ")


def word_count(text: str) -> int:
    return len(re.findall(r"\w+", text)) if isinstance(text, str) else 0


def slug_text(url: str) -> str:
    """Το path του URL ως «κείμενο» (greeklish slug) για θεματική ανίχνευση χωρίς λήψη."""
    from urllib.parse import unquote
    path = unquote(urlparse(url).path)
    return re.sub(r"[-_/]+", " ", path)


def load_themes(path: Path) -> dict[str, list[str]]:
    df = pd.read_csv(path)
    return {
        r.theme: [normalize(k.strip().replace("-", " ")) for k in str(r.keywords).split("|") if k.strip()]
        for r in df.itertuples()
    }


def detect_themes(text: str, themes: dict[str, list[str]]) -> list[str]:
    """Ένα θέμα ανιχνεύεται αν κάποια λέξη-κλειδί εμφανίζεται ως αρχή λέξης."""
    t = " " + normalize(text) + " "
    found = []
    for theme, kws in themes.items():
        for kw in kws:
            if re.search(r"(?<!\w)" + re.escape(kw), t):
                found.append(theme)
                break
    return found


# ─────────────────────────────────────────────
# URL
# ─────────────────────────────────────────────
def clean_url(u: str) -> str:
    u = u.strip().rstrip(".,;")
    p = urlparse(u)
    host = p.netloc
    if host in ("youtube", "www.youtube"):  # διόρθωση σπασμένου συνδέσμου στο αρχείο
        host = "youtu.be"
    q = [(k, v) for k, v in parse_qsl(p.query) if k not in TRACKING and not k.startswith("utm_")]
    return urlunparse((p.scheme, host, p.path, p.params, urlencode(q), ""))


def url_key(u: str) -> str:
    """Κλειδί για εντοπισμό διπλοεγγραφών (χωρίς www, /amp, τελική κάθετο, πεζά)."""
    p = urlparse(clean_url(u))
    host = p.netloc.lower().removeprefix("www.")
    path = re.sub(r"/amp/?$", "", p.path).rstrip("/")
    return f"{host}{path}" + (f"?{p.query}" if p.query else "")


def host_of(u: str) -> str:
    return urlparse(clean_url(u)).netloc.lower().removeprefix("www.")


def url_date(u: str) -> date | None:
    m = URL_DATE_RE.search(u)
    if m:
        try:
            return date(int(m[1]), int(m[2]), int(m[3]))
        except ValueError:
            return None
    return None


# ─────────────────────────────────────────────
# Ανάγνωση .docx
# ─────────────────────────────────────────────
def _is_bold(p) -> bool:
    runs = [r for r in p.runs if r.text.strip()]
    return bool(runs) and all(r.bold for r in runs)


def _infer_year(doc, default: int) -> int:
    for p in doc.paragraphs:
        m = DATE_HEAD_RE.match(p.text)
        if m and m[3]:
            return int(m[3])
    return default


def parse_docx(source, default_year: int = 2026) -> tuple[str, pd.DataFrame]:
    """Επιστρέφει (τίτλο, DataFrame δημοσιευμάτων). source: path ή bytes."""
    if isinstance(source, (bytes, bytearray)):
        source = BytesIO(source)
    doc = Document(source)
    year = _infer_year(doc, default_year)

    title = ""
    wave, wave_date = NO_WAVE, None
    cur_date: date | None = None
    mode = "article"  # article | interview | skip
    rows, order = [], 0

    for p in doc.paragraphs:
        text = p.text.strip()
        style = (p.style.name or "").lower()
        if not text:
            continue
        if style == "title" and not title:
            title = text
            continue

        if style.startswith("heading"):
            m = DATE_HEAD_RE.match(text)
            if m:
                cur_date = date(int(m[3] or year), int(m[2]), int(m[1]))
                mode = "article"
                if wave_date and (cur_date - wave_date).days > WAVE_WINDOW_DAYS:
                    wave, wave_date = NO_WAVE, None
            elif "domain" in text.lower():
                mode = "skip"
            elif "συνεντ" in normalize(text):
                mode, cur_date = "interview", None
            continue

        if mode == "skip":
            continue

        if _is_bold(p) and "τυπου" in normalize(text):
            wave = re.sub(r"\s+", " ", text)
            dm = DM_RE.search(text)
            wave_date = date(year, int(dm[2]), int(dm[1])) if dm else None
            cur_date = None
            mode = "article"
            continue

        urls = [h.address for h in p.hyperlinks if h.address and h.address.startswith("http")]
        urls += URL_RE.findall(text)
        seen, uniq = set(), []
        for u in urls:
            cu = clean_url(u)
            k = url_key(cu)
            if k not in seen:
                seen.add(k)
                uniq.append(cu)

        note = URL_RE.sub("", text).strip(" -–—()")
        base = dict(wave=wave, wave_date=wave_date, list_date=cur_date,
                    kind="Συνέντευξη" if mode == "interview" else "Δημοσίευμα")

        if uniq:
            for u in uniq:
                order += 1
                rows.append({**base, "order": order, "url": u, "url_key": url_key(u),
                             "host": host_of(u), "url_date": url_date(u), "note": note})
        elif cur_date or mode == "interview":
            # Έντυπο χωρίς σύνδεσμο, π.χ. «… (ΤΑ ΝΕΑ, 4 / 5 Ιουλίου 2026)»
            order += 1
            m = re.search(r"\(([^,()]+),", text)
            rows.append({**base, "order": order, "url": "", "url_key": f"print:{order}",
                         "host": "έντυπο", "url_date": None, "note": text,
                         "print_outlet": m[1].strip() if m else "Έντυπο"})

    df = pd.DataFrame(rows)
    if df.empty:
        return title, df
    if "print_outlet" not in df:
        df["print_outlet"] = None
    # Σύνδεσμοι συνεντεύξεων που υπάρχουν ήδη ως δημοσιεύματα → σήμανση ως συνέντευξη
    interview_keys = set(df.loc[df.kind == "Συνέντευξη", "url_key"])
    df.loc[df.url_key.isin(interview_keys), "kind"] = "Συνέντευξη"
    df["is_duplicate"] = df.duplicated("url_key", keep="first")
    for c in ("list_date", "url_date", "wave_date"):
        df[c] = pd.to_datetime(df[c])
    # Συνεντεύξεις χωρίς ημερομηνία: προσπάθεια από άλλη καταχώριση του ίδιου URL
    first_date = df.dropna(subset=["list_date"]).groupby("url_key")["list_date"].min()
    miss = df.list_date.isna()
    df.loc[miss, "list_date"] = df.loc[miss, "url_key"].map(first_date)
    return title, df


def resolve_dates(df: pd.DataFrame, prefer: str = "list") -> pd.DataFrame:
    """
    Τελική ημερομηνία ανά δημοσίευμα.
    prefer="list": ημ/νία αποδελτίωσης → ημ/νία από URL → ημ/νία μεταδεδομένων → ημ/νία Δ.Τ.
    prefer="pub" : ημ/νία δημοσίευσης (URL/μεταδεδομένα) → αποδελτίωση → Δ.Τ.
    """
    df = df.copy()
    meta = pd.to_datetime(df["pub_date_meta"], errors="coerce") if "pub_date_meta" in df else pd.Series(pd.NaT, index=df.index)
    if prefer == "pub":
        chain = [("URL", df.url_date), ("Μεταδεδομένα", meta), ("Αποδελτίωση", df.list_date), ("Δ.Τ. (κατά προσέγγιση)", df.wave_date)]
    else:
        chain = [("Αποδελτίωση", df.list_date), ("URL", df.url_date), ("Μεταδεδομένα", meta), ("Δ.Τ. (κατά προσέγγιση)", df.wave_date)]
    df["date"] = pd.NaT
    df["date_source"] = "—"
    for label, s in chain:
        fill = df["date"].isna() & s.notna()
        df.loc[fill, "date"] = s[fill]
        df.loc[fill, "date_source"] = label
    df["date"] = pd.to_datetime(df["date"])
    return df


def assign_period(d, start: date, end: date) -> str:
    if pd.isna(d):
        return "Άγνωστη ημ/νία"
    d = d.date()
    if d < start - timedelta(days=14):
        return "1. Προαναγγελία"
    if d < start:
        return "2. Αντίστροφη μέτρηση (−14 ημ.)"
    if d <= end:
        return "3. Διάρκεια φεστιβάλ"
    return "4. Απολογισμός"


def gini(values) -> float:
    import numpy as np
    x = np.sort(np.asarray(values, dtype=float))
    if len(x) == 0 or x.sum() == 0:
        return 0.0
    n = len(x)
    return float((2 * np.arange(1, n + 1) - n - 1).dot(x) / (n * x.sum()))


if __name__ == "__main__":
    import sys
    t, d = parse_docx(sys.argv[1] if len(sys.argv) > 1 else "data/Publicity5thCBF.docx")
    print(t, len(d), "εγγραφές,", int((~d.is_duplicate).sum()), "μοναδικές")
    print(d.groupby("wave").size())
